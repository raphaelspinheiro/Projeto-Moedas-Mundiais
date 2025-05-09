import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo

def criar_e_formatar_excel(nome_arquivo_entrada, nome_arquivo_saida):
    # Ler a planilha do Excel
    planilha = pd.read_excel(nome_arquivo_entrada)

    # Criar um novo arquivo Excel e adicionar uma planilha
    wb = Workbook()
    ws = wb.active

    # Transferir os dados da planilha para a planilha do novo arquivo Excel
    for r_idx, row in enumerate(dataframe_to_rows(planilha, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)

    # Criar uma tabela Excel formatada
    tab = Table(displayName="MinhaTabela", ref=ws.dimensions)
    style = TableStyleInfo(
        name="TableStyleMedium9", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=True)
    tab.tableStyleInfo = style
    ws.add_table(tab)

    # Salvar o novo arquivo Excel
    wb.save(nome_arquivo_saida)